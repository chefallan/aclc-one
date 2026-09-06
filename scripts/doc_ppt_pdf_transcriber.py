import sys
import os
import json
import base64
import tempfile

def find_python():
    candidates = [
        r"C:\Users\corte\Documents\projects NOT DELETE\lex memoria\venv\Scripts\python.exe",
        sys.executable
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return sys.executable

def extract_from_pptx(file_path):
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides_text = []
        images = []
        for idx, slide in enumerate(prs.slides):
            lines = []
            img_count = 0
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            lines.append(t)
                # Check for pictures (shape_type == 13 or has image)
                try:
                    if hasattr(shape, "image") and shape.image:
                        image_bytes = shape.image.blob
                        content_type = shape.image.content_type or 'image/jpeg'
                        if len(image_bytes) > 4000: # Skip tiny icons
                            img_count += 1
                            b64 = base64.b64encode(image_bytes).decode('utf-8')
                            images.append({
                                "title": f"Slide {idx + 1} - Figure {img_count}",
                                "caption": f"Extracted from Slide {idx + 1}",
                                "dataUrl": f"data:{content_type};base64,{b64}"
                            })
                except Exception:
                    pass

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    lines.append(f"Notes: {note}")
            
            slide_content = "\n".join(lines) if lines else "[Slide contains diagrams or images only]"
            slides_text.append(f"=== Slide {idx + 1} ===\n{slide_content}")
        return "\n\n".join(slides_text), images
    except Exception as e:
        return f"[PPTX extraction note: {str(e)}]", []

def extract_from_pdf(file_path):
    try:
        try:
            import fitz
        except ImportError:
            import pymupdf as fitz
        
        doc = fitz.open(file_path)
        pages_text = []
        images = []
        for idx, page in enumerate(doc):
            text = page.get_text("text").strip()
            image_list = page.get_images(full=True)
            if image_list:
                img_count = 0
                for img_info in image_list[:4]: # Cap per page
                    try:
                        xref = img_info[0]
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image.get("image", b"")
                        image_ext = base_image.get("ext", "jpeg")
                        if len(image_bytes) > 4000:
                            img_count += 1
                            b64 = base64.b64encode(image_bytes).decode('utf-8')
                            images.append({
                                "title": f"Page {idx + 1} - Figure {img_count}",
                                "caption": f"Extracted from Page {idx + 1}",
                                "dataUrl": f"data:image/{image_ext};base64,{b64}"
                            })
                    except Exception:
                        pass
                text += f"\n[Page contains {len(image_list)} figures/diagrams]"
            if not text:
                text = f"[Page {idx + 1}]"
            pages_text.append(f"=== Page {idx + 1} ===\n{text}")
        doc.close()
        return "\n\n".join(pages_text), images
    except Exception as e:
        return f"[PDF extraction note: {str(e)}]", []

def extract_from_docx(file_path):
    try:
        import zipfile
        import xml.etree.ElementTree as ET
        images = []
        text_lines = []
        with zipfile.ZipFile(file_path) as z:
            img_files = [f for f in z.filelist if f.filename.startswith("word/media/")]
            for idx, f in enumerate(img_files):
                img_bytes = z.read(f.filename)
                if len(img_bytes) > 3000:
                    ext = os.path.splitext(f.filename)[1].lower().replace('.', '')
                    mime = "image/png" if ext == "png" else ("image/gif" if ext == "gif" else "image/jpeg")
                    b64 = base64.b64encode(img_bytes).decode('utf-8')
                    images.append({
                        "title": f"Figure {idx + 1}",
                        "caption": f"Extracted from {os.path.basename(file_path)}",
                        "dataUrl": f"data:{mime};base64,{b64}"
                    })
            if "word/document.xml" in z.namelist():
                xml_content = z.read("word/document.xml")
                root = ET.fromstring(xml_content)
                for elem in root.iter():
                    if elem.tag.endswith('}p'):
                        p_text = "".join(node.text for node in elem.iter() if node.text)
                        if p_text.strip():
                            text_lines.append(p_text.strip())
        return "\n".join(text_lines), images
    except Exception as e:
        return f"[DOCX extraction note: {str(e)}]", []

def process_single_file(file_path):
    if not os.path.exists(file_path):
        return f"[File not found: {file_path}]", []
    
    ext = os.path.splitext(file_path)[1].lower()
    base_name = os.path.basename(file_path)
    images = []
    
    if ext in [".pptx", ".ppt"]:
        content, images = extract_from_pptx(file_path)
    elif ext == ".docx":
        content, images = extract_from_docx(file_path)
    elif ext == ".pdf":
        content, images = extract_from_pdf(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp"]:
        content = f"=== Image Reference: {base_name} ===\n[Visual diagram / educational diagram]"
        try:
            with open(file_path, "rb") as f:
                img_bytes = f.read()
                mime = "image/png" if ext == ".png" else "image/jpeg"
                b64 = base64.b64encode(img_bytes).decode('utf-8')
                images.append({
                    "title": base_name,
                    "caption": "Uploaded Image",
                    "dataUrl": f"data:{mime};base64,{b64}"
                })
        except Exception:
            pass
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
            
    return f"--- FILE: {base_name} ---\n{content}", images

def main():
    if len(sys.argv) < 2:
        print("Usage: doc_ppt_pdf_transcriber.py [--json] <file_path1> [file_path2 ...]", file=sys.stderr)
        sys.exit(1)
        
    want_json = "--json" in sys.argv
    file_paths = [arg for arg in sys.argv[1:] if arg != "--json"]
    
    results = []
    all_images = []
    for file_path in file_paths:
        text, imgs = process_single_file(file_path)
        results.append(text)
        all_images.extend(imgs)
        
    combined_text = "\n\n".join(results)
    
    if want_json:
        # Cap max images to 30 to keep IndexedDB lean
        payload = {
            "text": combined_text,
            "images": all_images[:30]
        }
        print(json.dumps(payload))
    else:
        print(combined_text)

if __name__ == "__main__":
    main()
